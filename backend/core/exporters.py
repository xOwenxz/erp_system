"""Exportação de DataFrames para Excel/CSV/JSON."""
import json
from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MEDIA_CSV = "text/csv; charset=utf-8"

_CARACTERES_INVALIDOS_ABA = set('[]:*?/\\')


def _nome_aba(nome: str) -> str:
    limpo = "".join(c for c in nome if c not in _CARACTERES_INVALIDOS_ABA).strip()
    return (limpo or "Planilha")[:31]


def to_xlsx(abas: dict[str, pd.DataFrame]) -> bytes:
    """Gera um .xlsx formatado (cabeçalho destacado, colunas ajustadas, painel congelado)."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for nome, df in abas.items():
            aba = _nome_aba(nome)
            df.to_excel(writer, index=False, sheet_name=aba)
            ws = writer.sheets[aba]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for idx, coluna in enumerate(df.columns, start=1):
                amostra = [len(str(v)) for v in df[coluna].head(300)]
                largura = max([len(str(coluna)), *amostra]) + 2
                ws.column_dimensions[get_column_letter(idx)].width = min(largura, 50)
            ws.freeze_panes = "A2"
    return buffer.getvalue()


def to_csv(df: pd.DataFrame) -> bytes:
    """CSV no padrão brasileiro (separador ';', decimal ',') com BOM p/ abrir no Excel."""
    return df.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")


def to_records(df: pd.DataFrame) -> list[dict]:
    """Converte para lista de dicts serializável (NaN -> null, datas -> ISO)."""
    return json.loads(df.to_json(orient="records", date_format="iso", force_ascii=False))
